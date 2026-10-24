import openpyxl
import datetime

wb = openpyxl.load_workbook('Data_Contoh_POS.xlsx', data_only=True)

def sql_val(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(int(v) if isinstance(v, float) and v.is_integer() else v)
    if isinstance(v, datetime.datetime):
        return f"'{v.isoformat()}'"
    if isinstance(v, datetime.date):
        return f"'{v.isoformat()}'"
    s = str(v).replace("'", "''")
    return f"'{s}'"

# 1. Cabang
cabang_rows = list(wb['cabang'].iter_rows(values_only=True))
cabang_header = cabang_rows[0]
cabang_data = cabang_rows[1:]

# 2. Profiles
profiles_rows = list(wb['profiles'].iter_rows(values_only=True))
profiles_header = profiles_rows[0][:6] # ignore cabang_kode
profiles_data = [r[:6] for r in profiles_rows[1:]]

# 3. Produk
produk_rows = list(wb['produk'].iter_rows(values_only=True))
produk_header = produk_rows[0]
produk_data = produk_rows[1:]

# 4. Riwayat harga modal
modal_rows = list(wb['riwayat_harga_modal'].iter_rows(values_only=True))
modal_header = modal_rows[0][:5] # ignore sku
modal_data = [r[:5] for r in modal_rows[1:]]

# 5. Stok cabang
stok_rows = list(wb['stok_cabang'].iter_rows(values_only=True))
stok_header = stok_rows[0][:6] # ignore helper columns
stok_data = [r[:6] for r in stok_rows[1:]]

# 6. Shift kasir
shift_rows = list(wb['shift_kasir'].iter_rows(values_only=True))
shift_header = shift_rows[0][:16] # ignore helper columns
shift_data = [r[:16] for r in shift_rows[1:]]

# 7. Transaksi
trans_rows = list(wb['transaksi'].iter_rows(values_only=True))
trans_header = trans_rows[0][:18] # ignore helper columns
trans_data = [r[:18] for r in trans_rows[1:]]

# 8. Transaksi item
item_rows = list(wb['transaksi_item'].iter_rows(values_only=True))
item_header = item_rows[0][:7] # ignore helper columns (modal_satuan is helper!)
item_data = [r[:7] for r in item_rows[1:]]

# 9. Audit log
audit_rows = list(wb['audit_log'].iter_rows(values_only=True))
audit_header = audit_rows[0][:8] # ignore helper columns
audit_data = [r[:8] for r in audit_rows[1:]]

print("Read all sheets successfully!")
print(f"Cabang: {len(cabang_data)}, Profiles: {len(profiles_data)}, Produk: {len(produk_data)}")
print(f"Modal: {len(modal_data)}, Stok: {len(stok_data)}, Shift: {len(shift_data)}")
print(f"Transaksi: {len(trans_data)}, Items: {len(item_data)}, Audit: {len(audit_data)}")
