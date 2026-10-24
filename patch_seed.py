import openpyxl
import datetime

wb = openpyxl.load_workbook('Data_Contoh_POS.xlsx', data_only=True)

def sql_val(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(int(v) if isinstance(v, float) and v.is_integer() else v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return f"'{v.isoformat()}'"
    s = str(v).replace("'", "''")
    return f"'{s}'"

# --- 1. RE-GENERATE SCHEMA.SQL WITH NULLABLE SHIFT TOTALS ---
with open('01_schema.sql', 'r', encoding='utf-8') as f:
    content = f.read()

content = content.replace("total_tunai INT NOT NULL DEFAULT 0,", "total_tunai INT DEFAULT 0,")
content = content.replace("total_qris INT NOT NULL DEFAULT 0,", "total_qris INT DEFAULT 0,")
content = content.replace("kas_seharusnya INT NOT NULL DEFAULT 0,", "kas_seharusnya INT,")

with open('01_schema.sql', 'w', encoding='utf-8') as f:
    f.write(content)

print("01_schema.sql updated!")

# --- 2. RE-GENERATE SEED.SQL WITH ALTER TABLE PREFIX ---
with open("02_seed.sql", "w", encoding="utf-8") as f:
    f.write("-- ==============================================================================\n")
    f.write("-- SEED DATA LUMBUNG KOPI DAN ROTI (DARI Data_Contoh_POS.xlsx)\n")
    f.write("-- ==============================================================================\n\n")

    # Pastikan kolom shift_kasir nullable untuk shift berstatus 'buka'
    f.write("-- Pastikan kolom shift_kasir nullable (sesuai wireframe halaman 3 untuk shift buka)\n")
    f.write("ALTER TABLE shift_kasir ALTER COLUMN total_tunai DROP NOT NULL;\n")
    f.write("ALTER TABLE shift_kasir ALTER COLUMN total_qris DROP NOT NULL;\n")
    f.write("ALTER TABLE shift_kasir ALTER COLUMN kas_seharusnya DROP NOT NULL;\n\n")

    # 1. Cabang
    f.write("-- 1. CABANG\n")
    for r in list(wb['cabang'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO cabang (id, kode, nama, alamat, aktif, urut_struk_tanggal, urut_struk_terakhir) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}, {sql_val(r[6])}) ON CONFLICT (id) DO UPDATE SET urut_struk_terakhir = EXCLUDED.urut_struk_terakhir;\n")
    f.write("\n")

    # 2. Auth Users & Profiles
    f.write("-- 2. AUTH USERS & PROFILES\n")
    prof_rows = list(wb['profiles'].iter_rows(values_only=True))[1:]
    for r in prof_rows:
        user_id = sql_val(r[0])
        nama = sql_val(r[1])
        email = sql_val(r[2])
        peran = sql_val(r[3])
        cabang_id = sql_val(r[4])
        aktif = sql_val(r[5])
        f.write(f"""INSERT INTO auth.users (id, instance_id, aud, role, email, encrypted_password, email_confirmed_at, raw_app_meta_data, raw_user_meta_data, created_at, updated_at)
VALUES ({user_id}, '00000000-0000-0000-0000-000000000000', 'authenticated', 'authenticated', {email}, crypt('Latihan#2026', gen_salt('bf')), NOW(), '{{"provider":"email","providers":["email"]}}', jsonb_build_object('nama_lengkap', {nama}, 'peran', {peran}), NOW(), NOW())
ON CONFLICT (id) DO NOTHING;

INSERT INTO profiles (id, nama_lengkap, email, peran, cabang_id, aktif)
VALUES ({user_id}, {nama}, {email}, {peran}, {cabang_id}, {aktif})
ON CONFLICT (id) DO UPDATE SET peran = EXCLUDED.peran, cabang_id = EXCLUDED.cabang_id, aktif = EXCLUDED.aktif;
""")
    f.write("\n")

    # 3. Produk
    f.write("-- 3. PRODUK\n")
    for r in list(wb['produk'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO produk (id, sku, nama, kategori, harga_jual, lacak_stok, aktif) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}, {sql_val(r[6])}) ON CONFLICT (id) DO UPDATE SET harga_jual = EXCLUDED.harga_jual, aktif = EXCLUDED.aktif;\n")
    f.write("\n")

    # 4. Riwayat Harga Modal
    f.write("-- 4. RIWAYAT HARGA MODAL\n")
    for r in list(wb['riwayat_harga_modal'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO riwayat_harga_modal (id, produk_id, harga_modal, berlaku_mulai, dibuat_oleh) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}) ON CONFLICT (id) DO NOTHING;\n")
    f.write("\n")

    # 5. Stok Cabang
    f.write("-- 5. STOK CABANG\n")
    for r in list(wb['stok_cabang'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO stok_cabang (id, cabang_id, produk_id, jumlah, stok_minimum, diubah_pada) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}) ON CONFLICT (id) DO UPDATE SET jumlah = EXCLUDED.jumlah, stok_minimum = EXCLUDED.stok_minimum;\n")
    f.write("\n")

    # 6. Shift Kasir
    f.write("-- 6. SHIFT KASIR\n")
    for r in list(wb['shift_kasir'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO shift_kasir (id, cabang_id, kasir_id, status, dibuka_pada, kas_awal, ditutup_pada, total_tunai, total_qris, kas_seharusnya, uang_fisik, selisih, catatan_tutup, diverifikasi_oleh, diverifikasi_pada, catatan_verifikasi) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}, {sql_val(r[6])}, {sql_val(r[7])}, {sql_val(r[8])}, {sql_val(r[9])}, {sql_val(r[10])}, {sql_val(r[11])}, {sql_val(r[12])}, {sql_val(r[13])}, {sql_val(r[14])}, {sql_val(r[15])}) ON CONFLICT (id) DO NOTHING;\n")
    f.write("\n")

    # 7. Transaksi
    f.write("-- 7. TRANSAKSI\n")
    for r in list(wb['transaksi'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO transaksi (id, nomor_struk, cabang_id, shift_id, kasir_id, status, subtotal, diskon, alasan_diskon, total, metode_bayar, uang_diterima, kembalian, dibuat_pada, dibayar_pada, void_oleh, void_pada, alasan_void) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}, {sql_val(r[6])}, {sql_val(r[7])}, {sql_val(r[8])}, {sql_val(r[9])}, {sql_val(r[10])}, {sql_val(r[11])}, {sql_val(r[12])}, {sql_val(r[13])}, {sql_val(r[14])}, {sql_val(r[15])}, {sql_val(r[16])}, {sql_val(r[17])}) ON CONFLICT (id) DO NOTHING;\n")
    f.write("\n")

    # 8. Transaksi Item
    f.write("-- 8. TRANSAKSI ITEM\n")
    for r in list(wb['transaksi_item'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO transaksi_item (id, transaksi_id, produk_id, nama_produk, harga_satuan, jumlah, subtotal) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}, {sql_val(r[6])}) ON CONFLICT (id) DO NOTHING;\n")
    f.write("\n")

    # 9. Audit Log
    f.write("-- 9. AUDIT LOG\n")
    for r in list(wb['audit_log'].iter_rows(values_only=True))[1:]:
        f.write(f"INSERT INTO audit_log (id, pelaku_id, aksi, tabel, data_id, nilai_sebelum, nilai_sesudah, dibuat_pada) VALUES ({sql_val(r[0])}, {sql_val(r[1])}, {sql_val(r[2])}, {sql_val(r[3])}, {sql_val(r[4])}, {sql_val(r[5])}::jsonb, {sql_val(r[6])}::jsonb, {sql_val(r[7])}) ON CONFLICT (id) DO NOTHING;\n")
    f.write("\n")

print("02_seed.sql regenerated successfully with nullable column alter!")
